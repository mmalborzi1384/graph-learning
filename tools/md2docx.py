"""تبدیل فایل‌های Markdown مستندات به Word (.docx) با جهت راست‌به‌چپ."""
import re, glob, os
from _paths import DOCS_DIR, ROOT
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def rtl_par(p):
    pPr = p._p.get_or_add_pPr()
    bidi = OxmlElement('w:bidi'); pPr.append(bidi)
    jc = OxmlElement('w:jc'); jc.set(qn('w:val'), 'right'); pPr.append(jc)

ISO_SPLIT = re.compile('(\u2066.*?\u2069)', re.S)

def _math_run(p, text, bold=False):
    """بلوک ریاضی: نویسه‌های ایزوله حفظ می‌شوند (ترتیب چپ‌به‌راست)، فونت ریاضی اعمال می‌شود."""
    r = p.add_run(text)
    r.bold = bold
    r.font.name = 'Cambria Math'
    r.font.rtl = False
    return r

def _emit(p, text, bold=False):
    for piece in ISO_SPLIT.split(text):
        if not piece:
            continue
        if piece.startswith('\u2066'):
            _math_run(p, piece, bold)
        else:
            r = p.add_run(piece); r.bold = bold

def add_runs(p, text):
    parts = re.split(r'(\*\*[^*]+\*\*)', text)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            _emit(p, part[2:-2], bold=True)
        else:
            _emit(p, re.sub(r'<[^>]+>', '', part))

def convert(md_path, out_path):
    doc = Document()
    st = doc.styles['Normal']; st.font.name = 'Tahoma'; st.font.size = Pt(11)
    st.element.rPr.rFonts.set(qn('w:cs'), 'Tahoma')
    lines = open(md_path, encoding='utf-8').read().splitlines()
    i = 0
    while i < len(lines):
        ln = lines[i].rstrip()
        if ln.startswith('|') and i + 1 < len(lines) and re.match(r'^\|[-| :]+\|$', lines[i+1].strip()):
            rows = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                if not re.match(r'^\|[-| :]+\|$', lines[i].strip()):
                    rows.append([c.strip() for c in lines[i].strip().strip('|').split('|')])
                i += 1
            ncol = max(len(r) for r in rows)
            t = doc.add_table(rows=len(rows), cols=ncol); t.style = 'Table Grid'
            for r, row in enumerate(rows):
                for c in range(ncol):
                    cell = t.cell(r, c); cell.text = ''
                    p = cell.paragraphs[0]; rtl_par(p)
                    add_runs(p, row[c] if c < len(row) else '')
                    if r == 0:
                        for run in p.runs: run.bold = True
            doc.add_paragraph()
            continue
        m = re.match(r'^(#{1,4})\s+(.*)', ln)
        if m:
            h = doc.add_heading(re.sub(r'\*\*', '', m.group(2)), level=min(len(m.group(1)), 4) if len(m.group(1)) > 1 else 0)
            rtl_par(h)
        elif ln.strip() == '---' or not ln.strip():
            pass
        elif re.match(r'^\s*[-*] ', ln):
            p = doc.add_paragraph(style='List Bullet'); rtl_par(p); add_runs(p, re.sub(r'^\s*[-*] ', '', ln))
        elif re.match(r'^\s*\d+\. ', ln):
            p = doc.add_paragraph(style='List Number'); rtl_par(p); add_runs(p, re.sub(r'^\s*\d+\. ', '', ln))
        elif ln.startswith('>'):
            p = doc.add_paragraph(); rtl_par(p); add_runs(p, ln.lstrip('> ')); 
            for r in p.runs: r.italic = True
        elif ln.startswith('```'):
            i += 1
            while i < len(lines) and not lines[i].startswith('```'):
                p = doc.add_paragraph(); r = p.add_run(lines[i]); r.font.name = 'Consolas'; i += 1
        else:
            p = doc.add_paragraph(); rtl_par(p); add_runs(p, ln)
        i += 1
    doc.save(out_path)

for md in sorted(glob.glob(str(DOCS_DIR / '*.md'))):
    out = os.path.splitext(md)[0] + '.docx'
    convert(md, out); print('OK', out)

